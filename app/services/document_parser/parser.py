import hashlib
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import structlog

logger = structlog.get_logger(__name__)


@dataclass
class ParsedSection:
    title: str | None
    content: str
    page_number: int | None
    clause_number: str | None
    section_type: str  # paragraph, clause, definition, schedule, heading, table


@dataclass
class ParsedDocument:
    raw_text: str
    sections: list[ParsedSection]
    page_count: int
    metadata: dict = field(default_factory=dict)
    tables: list[dict] = field(default_factory=list)


class DocumentParserService:
    """Parses PDF, DOCX, and scanned documents into structured text."""

    CLAUSE_NUMBER_PATTERN = re.compile(
        r'^(\d+(?:\.\d+)*)\s+',
        re.MULTILINE
    )
    SECTION_HEADING_PATTERN = re.compile(
        r'^(?:ARTICLE|SECTION|SCHEDULE|EXHIBIT|ANNEX|PART|CLAUSE)\s+[\dIVXLCDM]+',
        re.MULTILINE | re.IGNORECASE
    )
    DEFINITION_PATTERN = re.compile(
        r'"([^"]+)"\s+(?:means|shall mean|refers to|has the meaning)',
        re.IGNORECASE
    )

    async def parse(self, file_content: bytes, file_name: str, mime_type: str) -> ParsedDocument:
        logger.info("parsing_document", file_name=file_name, mime_type=mime_type)

        if mime_type == "application/pdf":
            return await self._parse_pdf(file_content, file_name)
        elif mime_type in (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/msword",
        ):
            return await self._parse_docx(file_content, file_name)
        elif mime_type.startswith("text/"):
            return await self._parse_text(file_content, file_name)
        elif mime_type.startswith("image/"):
            return await self._parse_image(file_content, file_name)
        else:
            raise ValueError(f"Unsupported file type: {mime_type}")

    async def _parse_pdf(self, content: bytes, file_name: str) -> ParsedDocument:
        import pdfplumber

        all_text_parts: list[str] = []
        sections: list[ParsedSection] = []
        tables: list[dict] = []
        page_count = 0

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            page_count = len(pdf.pages)

            for page_num, page in enumerate(pdf.pages, 1):
                page_text = page.extract_text() or ""

                if not page_text.strip():
                    page_text = await self._ocr_page(content, page_num)

                all_text_parts.append(page_text)

                page_tables = page.extract_tables()
                for table in page_tables:
                    if table:
                        tables.append({
                            "page": page_num,
                            "headers": table[0] if table else [],
                            "rows": table[1:] if len(table) > 1 else [],
                        })

                page_sections = self._extract_sections(page_text, page_num)
                sections.extend(page_sections)

        raw_text = "\n\n".join(all_text_parts)

        if not sections:
            sections = self._fallback_sectioning(raw_text, page_count)

        metadata = self._extract_pdf_metadata(content)

        return ParsedDocument(
            raw_text=raw_text,
            sections=sections,
            page_count=page_count,
            metadata=metadata,
            tables=tables,
        )

    async def _parse_docx(self, content: bytes, file_name: str) -> ParsedDocument:
        from docx import Document as DocxDocument
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = DocxDocument(io.BytesIO(content))
        all_text_parts: list[str] = []
        sections: list[ParsedSection] = []

        current_section_content: list[str] = []
        current_section_title: str | None = None
        current_clause_number: str | None = None

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            all_text_parts.append(text)

            is_heading = paragraph.style.name.startswith("Heading")
            is_bold_short = (
                paragraph.runs
                and all(r.bold for r in paragraph.runs if r.text.strip())
                and len(text) < 200
            )

            if is_heading or is_bold_short:
                if current_section_content:
                    sections.append(ParsedSection(
                        title=current_section_title,
                        content="\n".join(current_section_content),
                        page_number=None,
                        clause_number=current_clause_number,
                        section_type="clause" if current_clause_number else "paragraph",
                    ))
                    current_section_content = []

                current_section_title = text
                clause_match = self.CLAUSE_NUMBER_PATTERN.match(text)
                current_clause_number = clause_match.group(1) if clause_match else None
            else:
                current_section_content.append(text)

        if current_section_content:
            sections.append(ParsedSection(
                title=current_section_title,
                content="\n".join(current_section_content),
                page_number=None,
                clause_number=current_clause_number,
                section_type="clause" if current_clause_number else "paragraph",
            ))

        tables = []
        for table in doc.tables:
            rows_data = []
            for row in table.rows:
                rows_data.append([cell.text.strip() for cell in row.cells])
            if rows_data:
                tables.append({
                    "page": None,
                    "headers": rows_data[0],
                    "rows": rows_data[1:],
                })

        raw_text = "\n\n".join(all_text_parts)
        page_count = max(1, len(raw_text) // 3000)  # approximate

        return ParsedDocument(
            raw_text=raw_text,
            sections=sections if sections else self._fallback_sectioning(raw_text, page_count),
            page_count=page_count,
            metadata={"source_format": "docx"},
            tables=tables,
        )

    async def _parse_text(self, content: bytes, file_name: str) -> ParsedDocument:
        import chardet
        detected = chardet.detect(content)
        encoding = detected.get("encoding", "utf-8") or "utf-8"
        raw_text = content.decode(encoding, errors="replace")

        sections = self._extract_sections(raw_text, page_number=1)
        if not sections:
            sections = self._fallback_sectioning(raw_text, 1)

        return ParsedDocument(
            raw_text=raw_text,
            sections=sections,
            page_count=1,
            metadata={"source_format": "text", "encoding": encoding},
        )

    async def _parse_image(self, content: bytes, file_name: str) -> ParsedDocument:
        text = await self._ocr_image(content)
        sections = self._fallback_sectioning(text, 1)

        return ParsedDocument(
            raw_text=text,
            sections=sections,
            page_count=1,
            metadata={"source_format": "image", "ocr": True},
        )

    async def _ocr_page(self, pdf_content: bytes, page_num: int) -> str:
        try:
            from pdf2image import convert_from_bytes
            import pytesseract

            images = convert_from_bytes(
                pdf_content,
                first_page=page_num,
                last_page=page_num,
                dpi=300,
            )
            if images:
                text = pytesseract.image_to_string(images[0])
                return text.strip()
        except Exception as e:
            logger.warning("ocr_page_failed", page=page_num, error=str(e))
        return ""

    async def _ocr_image(self, content: bytes) -> str:
        try:
            import pytesseract
            from PIL import Image

            image = Image.open(io.BytesIO(content))
            text = pytesseract.image_to_string(image)
            return text.strip()
        except Exception as e:
            logger.warning("ocr_image_failed", error=str(e))
            return ""

    def _extract_sections(self, text: str, page_number: int | None) -> list[ParsedSection]:
        sections: list[ParsedSection] = []
        lines = text.split("\n")
        current_content: list[str] = []
        current_title: str | None = None
        current_clause: str | None = None

        for line in lines:
            stripped = line.strip()
            if not stripped:
                continue

            clause_match = self.CLAUSE_NUMBER_PATTERN.match(stripped)
            heading_match = self.SECTION_HEADING_PATTERN.match(stripped)

            if clause_match or heading_match or (stripped.isupper() and len(stripped) < 200):
                if current_content:
                    section_type = "clause" if current_clause else "paragraph"
                    if self.DEFINITION_PATTERN.search("\n".join(current_content)):
                        section_type = "definition"
                    sections.append(ParsedSection(
                        title=current_title,
                        content="\n".join(current_content),
                        page_number=page_number,
                        clause_number=current_clause,
                        section_type=section_type,
                    ))
                    current_content = []

                current_title = stripped
                current_clause = clause_match.group(1) if clause_match else None
            else:
                current_content.append(stripped)

        if current_content:
            sections.append(ParsedSection(
                title=current_title,
                content="\n".join(current_content),
                page_number=page_number,
                clause_number=current_clause,
                section_type="paragraph",
            ))

        return sections

    def _fallback_sectioning(self, text: str, page_count: int) -> list[ParsedSection]:
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        sections = []
        for i, para in enumerate(paragraphs):
            estimated_page = min(page_count, max(1, int((i / max(len(paragraphs), 1)) * page_count) + 1))
            sections.append(ParsedSection(
                title=None,
                content=para,
                page_number=estimated_page,
                clause_number=None,
                section_type="paragraph",
            ))
        return sections

    def _extract_pdf_metadata(self, content: bytes) -> dict:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(content))
            info = reader.metadata
            if info:
                return {
                    "title": info.get("/Title"),
                    "author": info.get("/Author"),
                    "subject": info.get("/Subject"),
                    "creator": info.get("/Creator"),
                    "source_format": "pdf",
                }
        except Exception:
            pass
        return {"source_format": "pdf"}

    @staticmethod
    def compute_hash(content: bytes) -> str:
        return hashlib.sha256(content).hexdigest()
